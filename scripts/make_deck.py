from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
import os

OUT = os.path.join(os.path.dirname(__file__), "../proposal/capabilities-deck.pptx")

# Colors
C_BG       = RGBColor(0x1e, 0x29, 0x3b)
C_DARK     = RGBColor(0x0f, 0x17, 0x2a)
C_ACCENT   = RGBColor(0x38, 0xbd, 0xf8)
C_GREEN    = RGBColor(0x4a, 0xde, 0x80)
C_WHITE    = RGBColor(0xf8, 0xfa, 0xfc)
C_MUTED    = RGBColor(0x94, 0xa3, 0xb8)
C_BORDER   = RGBColor(0x33, 0x41, 0x55)
C_LABEL    = RGBColor(0x64, 0x74, 0x8b)

W = Inches(13.33)
H = Inches(7.5)

prs = Presentation()
prs.slide_width  = W
prs.slide_height = H

blank = prs.slide_layouts[6]


def add_slide(prs):
    s = prs.slides.add_slide(blank)
    bg = s.background.fill
    bg.solid()
    bg.fore_color.rgb = C_BG
    return s


def txb(slide, text, x, y, w, h, size=18, bold=False, color=C_WHITE,
        align=PP_ALIGN.LEFT, wrap=True):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = wrap
    p  = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return tb


def rect(slide, x, y, w, h, fill=C_DARK, line=C_BORDER):
    shape = slide.shapes.add_shape(1, x, y, w, h)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = line
    shape.line.width = Emu(9144)
    return shape


def accent_bar(slide, x=Inches(0.8), y=Inches(2.2)):
    bar = slide.shapes.add_shape(1, x, y, Inches(0.5), Inches(0.04))
    bar.fill.solid()
    bar.fill.fore_color.rgb = C_ACCENT
    bar.line.fill.background()


def label(slide, text, x=Inches(0.8), y=Inches(0.4)):
    txb(slide, text, x, y, Inches(10), Inches(0.35),
        size=10, bold=True, color=C_LABEL)


# ── Slide 1: Title ──────────────────────────────────────────────────────────
s = add_slide(prs)
logo = rect(s, Inches(0.8), Inches(0.7), Inches(0.55), Inches(0.55), fill=C_ACCENT, line=C_ACCENT)
txb(s, "A", Inches(0.8), Inches(0.68), Inches(0.55), Inches(0.55),
    size=16, bold=True, color=C_DARK, align=PP_ALIGN.CENTER)
txb(s, "ACCENTURE", Inches(1.5), Inches(0.78), Inches(4), Inches(0.4),
    size=11, bold=True, color=C_LABEL)
accent_bar(s, Inches(0.8), Inches(1.7))
txb(s, "Inventory Dashboard Modernization", Inches(0.8), Inches(1.9),
    Inches(11), Inches(1.2), size=36, bold=True, color=C_WHITE)
txb(s, "Capabilities Presentation — Meridian Components\nRFP #MC-2026-0417  ·  April 2026",
    Inches(0.8), Inches(3.3), Inches(10), Inches(1),
    size=16, color=C_MUTED)

# ── Slide 2: The Situation ───────────────────────────────────────────────────
s = add_slide(prs)
label(s, "THE SITUATION")
txb(s, "A dashboard that mostly works.\nA team that can't fully rely on it.",
    Inches(0.8), Inches(0.9), Inches(11), Inches(1.4), size=26, bold=True)
bullets = [
    "Reports module has known defects — unresolved since the previous vendor's contract ended",
    "No automated test coverage — IT has blocked changes as a result",
    "Tokyo warehouse staff work in English-only views",
    "Restocking capability was requested but never delivered",
    "Previous vendor's documentation was minimal — architecture is undocumented",
]
for i, b in enumerate(bullets):
    y = Inches(2.5) + i * Inches(0.78)
    r = rect(s, Inches(0.8), y, Inches(11.5), Inches(0.65))
    txb(s, "·  " + b, Inches(1.0), y + Inches(0.1), Inches(11), Inches(0.5),
        size=13, color=C_MUTED)

# ── Slide 3: Our Understanding ───────────────────────────────────────────────
s = add_slide(prs)
label(s, "OUR UNDERSTANDING")
txb(s, "What Meridian actually needs", Inches(0.8), Inches(0.9),
    Inches(11), Inches(0.7), size=26, bold=True)
cards = [
    ("Operations team", "Needs a dashboard they can trust — fixed Reports, a working Restocking view, and screens their Tokyo colleagues can read."),
    ("IT team", "Won't approve changes without test coverage. Unblocking IT unblocks everything else. R3 is the prerequisite."),
    ("Procurement", "Needs a vendor who delivers on commitments, with predictable costs and milestone-based accountability."),
    ("The real ask", "Restore confidence. Prove that outside vendors can be trusted to finish what they start."),
]
for i, (title, body) in enumerate(cards):
    col = i % 2
    row = i // 2
    x = Inches(0.8) + col * Inches(6.0)
    y = Inches(1.9) + row * Inches(2.4)
    r = rect(s, x, y, Inches(5.7), Inches(2.1))
    txb(s, title, x + Inches(0.2), y + Inches(0.15), Inches(5.3), Inches(0.4),
        size=11, bold=True, color=C_MUTED)
    txb(s, body, x + Inches(0.2), y + Inches(0.55), Inches(5.3), Inches(1.4),
        size=13, color=C_WHITE)

# ── Slide 4: Scope Overview ──────────────────────────────────────────────────
s = add_slide(prs)
label(s, "SCOPE OF WORK")
txb(s, "Four required, three desired", Inches(0.8), Inches(0.9),
    Inches(11), Inches(0.6), size=26, bold=True)
required = [("R1", "Reports module remediation"), ("R2", "Restocking recommendations view"),
            ("R3", "Automated browser test coverage"), ("R4", "Architecture documentation")]
desired  = [("D1", "UI modernization"), ("D2", "Internationalization (Tokyo priority)"),
            ("D3", "Dark mode")]
txb(s, "REQUIRED", Inches(0.8), Inches(1.75), Inches(5), Inches(0.3),
    size=10, bold=True, color=C_ACCENT)
for i, (tag, text) in enumerate(required):
    y = Inches(2.15) + i * Inches(0.9)
    r = rect(s, Inches(0.8), y, Inches(5.7), Inches(0.75))
    txb(s, tag, Inches(1.0), y + Inches(0.15), Inches(0.6), Inches(0.45),
        size=12, bold=True, color=C_ACCENT)
    txb(s, text, Inches(1.65), y + Inches(0.18), Inches(4.5), Inches(0.45),
        size=13, color=C_WHITE)
txb(s, "DESIRED", Inches(7.0), Inches(1.75), Inches(5), Inches(0.3),
    size=10, bold=True, color=C_GREEN)
for i, (tag, text) in enumerate(desired):
    y = Inches(2.15) + i * Inches(0.9)
    r = rect(s, Inches(7.0), y, Inches(5.5), Inches(0.75))
    txb(s, tag, Inches(7.2), y + Inches(0.15), Inches(0.6), Inches(0.45),
        size=12, bold=True, color=C_GREEN)
    txb(s, text, Inches(7.85), y + Inches(0.18), Inches(4.4), Inches(0.45),
        size=13, color=C_WHITE)

# ── Slide 5: R1 ──────────────────────────────────────────────────────────────
s = add_slide(prs)
label(s, "TECHNICAL APPROACH — R1")
txb(s, "Reports Module Remediation", Inches(0.8), Inches(0.9),
    Inches(11), Inches(0.7), size=26, bold=True)
points = [
    "Start with a structured discovery audit — no issue log exists, so we identify defects before we fix them",
    "Known areas: filter wiring incomplete, Options API patterns inconsistent with the rest of the codebase",
    "Audit findings shared with Meridian operations team before remediation begins",
    "All defects identified in discovery are in scope — no surprises",
    "i18n gaps in Reports addressed here; remaining views covered under D2",
]
for i, p in enumerate(points):
    y = Inches(2.0) + i * Inches(0.92)
    r = rect(s, Inches(0.8), y, Inches(11.5), Inches(0.78))
    txb(s, "·  " + p, Inches(1.0), y + Inches(0.14), Inches(11), Inches(0.55),
        size=13, color=C_MUTED)

# ── Slide 6: R2 ──────────────────────────────────────────────────────────────
s = add_slide(prs)
label(s, "TECHNICAL APPROACH — R2")
txb(s, "Restocking Recommendations", Inches(0.8), Inches(0.9),
    Inches(11), Inches(0.7), size=26, bold=True)
points = [
    "New first-class view — consistent with existing navigation and layout patterns",
    "Uses data already in the system: stock levels, demand forecasts, historical purchase order costs",
    "Operator supplies a budget ceiling; system returns a prioritized purchase order list",
    "Output: item, suggested quantity, estimated unit cost, line total — ranked by demand urgency",
    "New backend endpoint — clean, follows existing API patterns",
]
for i, p in enumerate(points):
    y = Inches(2.0) + i * Inches(0.92)
    r = rect(s, Inches(0.8), y, Inches(11.5), Inches(0.78))
    txb(s, "·  " + p, Inches(1.0), y + Inches(0.14), Inches(11), Inches(0.55),
        size=13, color=C_MUTED)

# ── Slide 7: R3 & R4 ─────────────────────────────────────────────────────────
s = add_slide(prs)
label(s, "TECHNICAL APPROACH — R3 & R4")
txb(s, "Testing & Documentation", Inches(0.8), Inches(0.9),
    Inches(11), Inches(0.7), size=26, bold=True)
for col, (tag, title, body) in enumerate([
    ("R3", "Browser Testing",
     "End-to-end Playwright tests covering inventory by warehouse, order filtering, and the Restocking view. Structured for IT sign-off. Tests ship alongside features — not at the end."),
    ("R4", "Architecture Docs",
     "Self-contained HTML diagram: Vue frontend → FastAPI backend → JSON data layer. All views mapped to API endpoints. Written for Meridian IT, produced at engagement start."),
]):
    x = Inches(0.8) + col * Inches(6.2)
    r = rect(s, x, Inches(2.0), Inches(5.9), Inches(4.5))
    txb(s, tag, x + Inches(0.25), Inches(2.2), Inches(1), Inches(0.45),
        size=14, bold=True, color=C_ACCENT)
    txb(s, title, x + Inches(0.25), Inches(2.75), Inches(5.4), Inches(0.5),
        size=16, bold=True, color=C_WHITE)
    txb(s, body, x + Inches(0.25), Inches(3.4), Inches(5.4), Inches(2.8),
        size=13, color=C_MUTED)

# ── Slide 8: Timeline ────────────────────────────────────────────────────────
s = add_slide(prs)
label(s, "TIMELINE")
txb(s, "Six weeks, three phases", Inches(0.8), Inches(0.9),
    Inches(11), Inches(0.6), size=26, bold=True)
phases = [
    ("Weeks 1–2", "Phase 1\nFoundation", "Architecture docs (R4)\nReports audit\nTest infrastructure"),
    ("Weeks 3–4", "Phase 2\nRemediation", "Reports fixes (R1)\nInitial tests (R3)\ni18n gaps in Reports"),
    ("Weeks 5–6", "Phase 3\nNew Build", "Restocking view (R2)\nFull test suite (R3)\nUI refresh + i18n"),
    ("Post Wk 6",  "Stretch",            "Dark mode (D3)\nIf elected at\ncontract signature"),
]
for i, (week, title, items) in enumerate(phases):
    x = Inches(0.8) + i * Inches(3.1)
    r = rect(s, x, Inches(1.9), Inches(2.9), Inches(4.5))
    txb(s, week,  x + Inches(0.2), Inches(2.05), Inches(2.5), Inches(0.35),
        size=10, bold=True, color=C_LABEL)
    txb(s, title, x + Inches(0.2), Inches(2.45), Inches(2.5), Inches(0.8),
        size=14, bold=True, color=C_WHITE)
    txb(s, items, x + Inches(0.2), Inches(3.35), Inches(2.5), Inches(2.5),
        size=12, color=C_MUTED)
txb(s, "Foundation first — every change made safely, IT can approve at each milestone.",
    Inches(0.8), Inches(6.7), Inches(11), Inches(0.4), size=12, color=C_LABEL)

# ── Slide 9: Pricing ─────────────────────────────────────────────────────────
s = add_slide(prs)
label(s, "PRICING")
txb(s, "Fixed-fee. No surprises.", Inches(0.8), Inches(0.9),
    Inches(11), Inches(0.6), size=26, bold=True)
rows = [
    ("R4 — Architecture documentation",   "$4,000"),
    ("R1 — Reports module remediation",   "$12,000"),
    ("R3 — Automated browser testing",    "$8,000"),
    ("R2 — Restocking recommendations",   "$18,000"),
    ("D1 + D2 — UI modernization & i18n", "$11,000"),
]
for i, (lbl, amt) in enumerate(rows):
    y = Inches(1.85) + i * Inches(0.82)
    r = rect(s, Inches(0.8), y, Inches(11.5), Inches(0.68))
    txb(s, lbl, Inches(1.0), y + Inches(0.15), Inches(8), Inches(0.45),
        size=13, color=C_MUTED)
    txb(s, amt, Inches(9.5), y + Inches(0.12), Inches(2.5), Inches(0.45),
        size=14, bold=True, color=C_ACCENT, align=PP_ALIGN.RIGHT)
# Total
y = Inches(6.0)
r = rect(s, Inches(0.8), y, Inches(11.5), Inches(0.78), fill=RGBColor(0x0f,0x17,0x2a))
txb(s, "Total (R1–R4, D1–D2)", Inches(1.0), y + Inches(0.17), Inches(8), Inches(0.5),
    size=15, bold=True, color=C_WHITE)
txb(s, "$53,000", Inches(9.5), y + Inches(0.14), Inches(2.5), Inches(0.5),
    size=18, bold=True, color=C_GREEN, align=PP_ALIGN.RIGHT)

# ── Slide 10: Why Us ─────────────────────────────────────────────────────────
s = add_slide(prs)
label(s, "WHY US")
txb(s, "We finish what we start.", Inches(0.8), Inches(0.9),
    Inches(11), Inches(0.6), size=26, bold=True)
cards = [
    ("We audit before we fix",
     "Discovery findings shared with your team before work begins. No hidden scope, no end-of-project surprises."),
    ("Tests ship with features",
     "Every deliverable includes coverage. IT can review and approve at each milestone — not just at the end."),
    ("Documentation is a deliverable",
     "Architecture docs written for your IT team. You won't need another vendor to explain what we built."),
    ("Milestone accountability",
     "Fixed fee, phased payments. You pay for results at each milestone, not hours on a timesheet."),
]
for i, (title, body) in enumerate(cards):
    col = i % 2
    row = i // 2
    x = Inches(0.8) + col * Inches(6.2)
    y = Inches(1.9) + row * Inches(2.5)
    r = rect(s, x, y, Inches(5.9), Inches(2.2))
    txb(s, title, x + Inches(0.25), y + Inches(0.2), Inches(5.4), Inches(0.45),
        size=14, bold=True, color=C_WHITE)
    txb(s, body, x + Inches(0.25), y + Inches(0.75), Inches(5.4), Inches(1.3),
        size=13, color=C_MUTED)

# ── Slide 11: Close ──────────────────────────────────────────────────────────
s = add_slide(prs)
logo2 = rect(s, Inches(0.8), Inches(0.7), Inches(0.55), Inches(0.55), fill=C_ACCENT, line=C_ACCENT)
txb(s, "A", Inches(0.8), Inches(0.68), Inches(0.55), Inches(0.55),
    size=16, bold=True, color=C_DARK, align=PP_ALIGN.CENTER)
txb(s, "ACCENTURE", Inches(1.5), Inches(0.78), Inches(4), Inches(0.4),
    size=11, bold=True, color=C_LABEL)
accent_bar(s, Inches(0.8), Inches(1.7))
txb(s, "Meridian's operations team deserves\na dashboard they can trust.",
    Inches(0.8), Inches(1.9), Inches(11), Inches(1.6), size=28, bold=True, color=C_WHITE)
txb(s, "That's what we're here to deliver.",
    Inches(0.8), Inches(3.7), Inches(11), Inches(0.6), size=18, color=C_MUTED)
cta = rect(s, Inches(0.8), Inches(4.7), Inches(11.5), Inches(1.2))
txb(s, "Questions or to schedule a follow-up presentation",
    Inches(0.8), Inches(4.9), Inches(11.5), Inches(0.4),
    size=13, color=C_MUTED, align=PP_ALIGN.CENTER)
txb(s, "pawel.chmielewski@accenture.com",
    Inches(0.8), Inches(5.35), Inches(11.5), Inches(0.4),
    size=15, bold=True, color=C_ACCENT, align=PP_ALIGN.CENTER)

prs.save(OUT)
print(f"Saved: {OUT}")
